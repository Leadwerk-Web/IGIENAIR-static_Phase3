#!/usr/bin/env python3
"""Erstellt die PageSpeed-Ergebnis-Excel: Alte Site (15.07. nachmittags) vs.
neue Site (15.07. abends, nach Go-Live) fuer Desktop und Mobil, inkl.
eingebetteter Screenshots der pagespeed.web.dev-Ergebnisse."""

import os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(BASE, "screenshots")
OUT = os.path.join(BASE, "Igienair_PageSpeed_Scores.xlsx")

# Pro Seite: key, Name, URL, dann je Messung (alt desktop, neu desktop,
# alt mobil, neu mobil) ein Tupel (perf, fcp, lcp, tbt, cls, si).
PAGES = [
    ("home", "Startseite", "https://igienair.de/",
     (91, "0,5 s", "0,6 s", "40 ms", "0.143", "1,9 s"),
     (96, "0,4 s", "0,7 s", "90 ms", "0.001", "1,9 s"),
     (46, "2,1 s", "7,9 s", "160 ms", "0.351", "9,3 s"),
     (87, "2,0 s", "3,5 s", "50 ms", "0", "4,0 s")),
    ("unternehmen", "Unternehmen", "https://igienair.de/unternehmen/",
     (90, "0,6 s", "0,8 s", "90 ms", "0.096", "2,7 s"),
     (97, "0,4 s", "0,8 s", "0 ms", "0.098", "0,5 s"),
     (42, "1,9 s", "6,3 s", "360 ms", "0.353", "8,4 s"),
     (87, "1,5 s", "4,0 s", "0 ms", "0", "2,5 s")),
    ("zertifizierungen", "Zertifizierungen", "https://igienair.de/unternehmen/zertifizierungen/",
     (93, "0,5 s", "0,8 s", "100 ms", "0.101", "1,7 s"),
     (100, "0,3 s", "0,7 s", "0 ms", "0", "0,5 s"),
     (44, "2,1 s", "8,4 s", "280 ms", "0.344", "6,9 s"),
     (85, "1,7 s", "4,1 s", "0 ms", "0", "3,1 s")),
    ("hygieneinspektion", "Hygieneinspektion VDI 6022", "https://igienair.de/hygieneinspektion-vdi-6022/",
     (94, "0,5 s", "0,8 s", "30 ms", "0.099", "1,8 s"),
     (98, "0,4 s", "1,1 s", "0 ms", "0.001", "0,6 s"),
     (48, "2,1 s", "8,9 s", "100 ms", "0.34", "8,8 s"),
     (74, "1,8 s", "6,1 s", "20 ms", "0", "4,4 s")),
    ("lueftungsreinigung", "Lüftungsreinigung", "https://igienair.de/lueftungsreinigung/",
     (73, "0,5 s", "0,8 s", "360 ms", "0.117", "2,9 s"),
     (91, "0,4 s", "1,9 s", "0 ms", "0.002", "0,8 s"),
     (50, "2,1 s", "8,6 s", "120 ms", "0.343", "6,0 s"),
     (71, "1,7 s", "12,8 s", "0 ms", "0.001", "4,7 s")),
    ("luftkanalreinigung", "Luftkanalreinigung", "https://igienair.de/luftkanalreinigung/",
     (93, "0,5 s", "0,8 s", "100 ms", "0.101", "1,9 s"),
     (84, "0,4 s", "2,2 s", "180 ms", "0.002", "0,9 s"),
     (48, "2,1 s", "9,5 s", "170 ms", "0.346", "6,7 s"),
     (69, "1,7 s", "12,6 s", "10 ms", "0", "5,7 s")),
    ("kuehlturmreinigung", "Kühlturmreinigung", "https://igienair.de/kuehlturmreinigung/",
     (88, "0,5 s", "1,1 s", "150 ms", "0.106", "1,9 s"),
     (91, "0,4 s", "1,9 s", "0 ms", "0.001", "0,9 s"),
     (46, "2,1 s", "9,3 s", "230 ms", "0.347", "6,8 s"),
     (74, "1,7 s", "10,8 s", "0 ms", "0.001", "2,7 s")),
    ("branchen", "Branchen", "https://igienair.de/branchen/",
     (88, "0,4 s", "0,7 s", "200 ms", "0.105", "1,7 s"),
     (100, "0,4 s", "0,6 s", "0 ms", "0", "0,5 s"),
     (39, "1,8 s", "6,9 s", "430 ms", "0.355", "9,1 s"),
     (91, "1,7 s", "3,3 s", "0 ms", "0", "2,8 s")),
    ("angebot", "Angebot anfordern", "https://igienair.de/kontakt/angebot-anfordern/",
     (72, "0,5 s", "1,2 s", "370 ms", "0.11", "2,6 s"),
     (99, "0,4 s", "1,0 s", "40 ms", "0.001", "0,5 s"),
     (52, "2,6 s", "8,9 s", "100 ms", "0.3", "5,4 s"),
     (87, "1,7 s", "3,9 s", "60 ms", "0.003", "1,7 s")),
    ("glossar", "Glossar", "https://igienair.de/glossar/",
     (88, "0,7 s", "1,2 s", "60 ms", "0.096", "2,5 s"),
     (100, "0,4 s", "0,6 s", "0 ms", "0", "0,5 s"),
     (52, "4,1 s", "5,6 s", "140 ms", "0.245", "6,0 s"),
     (93, "1,7 s", "3,1 s", "0 ms", "0.001", "1,7 s")),
]

GREEN = PatternFill("solid", fgColor="C6EFCE")
ORANGE = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
HEADER_FILL = PatternFill("solid", fgColor="073975")
HEADER_FONT = Font(bold=True, color="FFFFFF")
DELTA_UP = Font(bold=True, color="1E7E34")
DELTA_DOWN = Font(bold=True, color="C82333")


def score_fill(score):
    if score >= 90:
        return GREEN
    if score >= 50:
        return ORANGE
    return RED


wb = Workbook()


def build_scores_sheet(ws, old_idx, new_idx):
    headers = [
        "Seite", "URL",
        "Perf. ALT", "Perf. NEU", "Δ Perf.",
        "FCP alt", "FCP neu", "LCP alt", "LCP neu",
        "TBT alt", "TBT neu", "CLS alt", "CLS neu",
        "Speed Index alt", "Speed Index neu",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for page in PAGES:
        name, url = page[1], page[2]
        old, new = page[old_idx], page[new_idx]
        delta = new[0] - old[0]
        ws.append([
            name, url,
            old[0], new[0], ("+" if delta >= 0 else "") + str(delta),
            old[1], new[1], old[2], new[2],
            old[3], new[3], old[4], new[4],
            old[5], new[5],
        ])

    for r in range(2, len(PAGES) + 2):
        for col in (3, 4):
            cell = ws.cell(row=r, column=col)
            cell.fill = score_fill(int(cell.value))
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        dcell = ws.cell(row=r, column=5)
        dcell.font = DELTA_UP if not str(dcell.value).startswith("-") else DELTA_DOWN
        dcell.alignment = Alignment(horizontal="center")
        for col in range(6, 16):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center")

    widths = [28, 46, 10, 10, 9] + [11] * 10
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"

    meta = ws.cell(row=len(PAGES) + 3, column=1,
                   value="ALT = alte Website (15.07.2026 nachmittags), NEU = neue Website nach Go-Live "
                         "(15.07.2026 abends). Quelle: pagespeed.web.dev (Lighthouse).")
    meta.font = Font(italic=True, size=10)


ws_d = wb.active
ws_d.title = "Scores Desktop"
build_scores_sheet(ws_d, 3, 4)

ws_m = wb.create_sheet("Scores Mobil")
build_scores_sheet(ws_m, 5, 6)

# ---------- Screenshot-Sheets ----------
IMG_W = 900
ROW_STEP = 36


def add_shots(sheet_name, prefix):
    s = wb.create_sheet(sheet_name)
    s.column_dimensions["A"].width = 120
    row = 1
    for page in PAGES:
        key, name, url = page[0], page[1], page[2]
        title = s.cell(row=row, column=1, value=f"{name} – {url}")
        title.font = Font(bold=True, size=12)
        path = os.path.join(SHOTS, f"{prefix}{key}.png")
        if os.path.exists(path):
            img = XLImage(path)
            scale = IMG_W / img.width
            img.width = IMG_W
            img.height = int(img.height * scale)
            s.add_image(img, f"A{row + 1}")
        row += ROW_STEP


add_shots("SS Desktop ALT", "psi_")
add_shots("SS Desktop NEU", "psi_new_")
add_shots("SS Mobil ALT", "psi_mobile_")
add_shots("SS Mobil NEU", "psi_new_mobile_")

wb.save(OUT)
print("geschrieben:", OUT)
